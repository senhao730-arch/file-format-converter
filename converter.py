"""Local conversion engines. No uploads; source files are never modified."""
from __future__ import annotations

import io
import csv
import os
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

IMAGE_INPUTS = {'png', 'jpg', 'jpeg', 'webp', 'bmp', 'tif', 'tiff', 'gif', 'ico', 'avif', 'heic', 'heif'}
IMAGE_OUTPUTS = ('png', 'jpg', 'webp', 'bmp', 'tiff', 'gif', 'ico', 'avif')
VIDEO_INPUTS = {'mp4', 'mov', 'mkv', 'avi', 'webm', 'wmv', 'flv', 'mpeg', 'mpg', 'm4v', 'ts', 'mts', 'm2ts', '3gp', 'vob', 'ogv'}
VIDEO_OUTPUTS = ('mp4', 'mov', 'mkv', 'avi', 'webm', 'wmv', 'flv', 'mpg', 'm4v', 'ts', 'gif')
WORD_INPUTS = {'doc', 'docx', 'rtf', 'odt'}
PPT_INPUTS = {'ppt', 'pptx', 'odp', 'pps', 'ppsx'}
CATEGORIES = ('Word / PDF', 'PPT / PDF', '图片', 'GIF / 视频', '视频')
CREATE_NO_WINDOW = getattr(subprocess, 'CREATE_NO_WINDOW', 0)


class ConversionError(Exception):
    pass


class Cancelled(ConversionError):
    pass


@dataclass
class Options:
    quality: int = 90
    dpi: int = 150
    pdf_word_mode: str = 'editable'
    video_quality: str = 'balanced'
    timeout: int = 3600
    office_backend: str = 'auto'
    width: int = 0
    height: int = 0
    resize_mode: str = 'contain'
    sharpen: float = 0
    denoise: float = 0
    watermark_boxes: tuple = ()
    watermark_start: float = 0
    watermark_end: float = 0
    crop_box: tuple = ()


def targets_for(source: Path | str, category: str) -> tuple[str, ...]:
    ext = Path(source).suffix.lower().lstrip('.')
    if category == CATEGORIES[0]:
        return ('pdf',) if ext in WORD_INPUTS else (('docx',) if ext == 'pdf' else ())
    if category == CATEGORIES[1]:
        return ('pdf',) if ext in PPT_INPUTS else (('pptx',) if ext == 'pdf' else ())
    if category == CATEGORIES[2]:
        return IMAGE_OUTPUTS if ext in IMAGE_INPUTS else ()
    if category == CATEGORIES[3]:
        return VIDEO_OUTPUTS if ext == 'gif' or ext in VIDEO_INPUTS else ()
    if category == CATEGORIES[4]:
        return VIDEO_OUTPUTS if ext in VIDEO_INPUTS else ()
    return ()


def check_cancel(cancel: threading.Event) -> None:
    if cancel.is_set():
        raise Cancelled('已取消')


def ffmpeg_path() -> str:
    configured = os.environ.get('FFMPEG_BINARY')
    if configured and Path(configured).is_file():
        return configured
    local = Path(sys.executable if getattr(sys, 'frozen', False) else __file__).parent / 'bin' / 'ffmpeg.exe'
    if local.is_file():
        return str(local)
    system = shutil.which('ffmpeg')
    if system:
        return system
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception as exc:
        raise ConversionError('未找到 FFmpeg。请安装依赖，或把 ffmpeg.exe 放入程序目录的 bin 文件夹。') from exc


def libreoffice_path() -> str | None:
    configured = os.environ.get('LIBREOFFICE_BINARY')
    if configured and Path(configured).is_file():
        return configured
    for name in ('soffice', 'libreoffice'):
        if found := shutil.which(name):
            return found
    for base in (os.environ.get('PROGRAMFILES'), os.environ.get('PROGRAMFILES(X86)')):
        if base:
            path = Path(base) / 'LibreOffice' / 'program' / 'soffice.exe'
            if path.is_file():
                return str(path)
    return None


def publish(staged: Path, directory: Path, stem: str, extension: str) -> Path:
    """Create a fresh filename exclusively, including when another process races us."""
    directory.mkdir(parents=True, exist_ok=True)
    for number in range(100000):
        suffix = '' if number == 0 else f' ({number})'
        destination = directory / f'{stem}{suffix}.{extension}'
        try:
            with destination.open('xb') as output:
                try:
                    with staged.open('rb') as source:
                        shutil.copyfileobj(source, output, 1024 * 1024)
                except BaseException:
                    output.close()
                    destination.unlink(missing_ok=True)
                    raise
        except FileExistsError:
            continue
        return destination
    raise ConversionError('同名结果文件过多，请选择其他输出目录。')


def _image(source: Path, target: Path, options: Options, cancel: threading.Event,
           progress: Callable[[float, str], None], notes: list[str]) -> None:
    from PIL import Image, ImageOps, ImageSequence
    from media import adjust_image, media_notes
    notes.extend(media_notes(options))
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except ImportError:
        pass
    Image.init()
    format_name = {'jpg': 'JPEG', 'tiff': 'TIFF'}.get(target.suffix[1:], target.suffix[1:].upper())
    if format_name not in Image.SAVE:
        raise ConversionError(f'当前图像引擎不支持写入 {format_name}。')
    with Image.open(source) as original:
        count = getattr(original, 'n_frames', 1)
        animated_target = format_name in {'GIF', 'WEBP', 'PNG', 'TIFF', 'AVIF'}
        frames, durations = [], []
        loop = original.info.get('loop', 0)
        for index, frame in enumerate(ImageSequence.Iterator(original)):
            check_cancel(cancel)
            if index > 0 and not animated_target:
                notes.append('目标格式为静态图像，仅保存第一帧。')
                break
            durations.append(frame.info.get('duration', 100))
            converted = ImageOps.exif_transpose(frame).convert('RGBA')
            converted = adjust_image(converted, options)
            if format_name in {'JPEG', 'BMP'}:
                canvas = Image.new('RGB', converted.size, 'white')
                canvas.paste(converted, mask=converted.getchannel('A'))
                converted = canvas
            elif format_name == 'TIFF':
                converted = converted.convert('RGB') if 'A' not in frame.getbands() else converted
            elif format_name == 'ICO':
                converted.thumbnail((256, 256), Image.Resampling.LANCZOS)
                notes.append('ICO 最大边已限制为 256 像素。')
            frames.append(converted)
            progress((index + 1) / count * 0.85, f'读取图像 {index + 1}/{count}')
        kwargs = {}
        if format_name in {'JPEG', 'WEBP', 'AVIF'}:
            kwargs['quality'] = options.quality
        if format_name == 'JPEG':
            kwargs['subsampling'] = 0
        if format_name == 'TIFF':
            kwargs['compression'] = 'tiff_deflate'
        if len(frames) > 1:
            kwargs.update(save_all=True, append_images=frames[1:])
            if format_name != 'TIFF':
                kwargs.update(duration=durations, loop=loop)
            if format_name == 'GIF':
                kwargs.update(disposal=2, optimize=False)
        if profile := original.info.get('icc_profile'):
            if format_name in {'JPEG', 'PNG', 'WEBP', 'TIFF', 'AVIF'}:
                kwargs['icc_profile'] = profile
        check_cancel(cancel)
        frames[0].save(target, format=format_name, **kwargs)
        if format_name == 'GIF':
            notes.append('GIF 最多支持 256 色，部分色彩会量化。')


def _open_pdf(source: Path):
    import pymupdf
    document = pymupdf.open(source)
    if document.needs_pass:
        document.close()
        raise ConversionError('PDF 有密码保护，请先解密后再转换。')
    if not document.page_count:
        document.close()
        raise ConversionError('PDF 没有可转换的页面。')
    return document


def _pdf_to_ppt(source, target, options, cancel, progress, notes):
    from pptx import Presentation
    from pptx.util import Inches
    with _open_pdf(source) as document:
        presentation = Presentation()
        # PowerPoint accepts slide dimensions between 1 and 56 inches.
        width, height = document[0].rect.width / 72, document[0].rect.height / 72
        factor = min(1, 56 / max(width, height))
        width, height = max(1, width * factor), max(1, height * factor)
        presentation.slide_width = Inches(width)
        presentation.slide_height = Inches(height)
        for index, page in enumerate(document):
            check_cancel(cancel)
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            pixmap = page.get_pixmap(dpi=options.dpi, alpha=False)
            ratio = min(presentation.slide_width / pixmap.width, presentation.slide_height / pixmap.height)
            pw, ph = int(pixmap.width * ratio), int(pixmap.height * ratio)
            slide.shapes.add_picture(io.BytesIO(pixmap.tobytes('png')),
                                     int((presentation.slide_width - pw) / 2),
                                     int((presentation.slide_height - ph) / 2), pw, ph)
            progress((index + 1) / len(document) * 0.95, f'转换 PDF 页 {index + 1}/{len(document)}')
        presentation.save(target)
    notes.append('每页 PDF 已作为整页图片放入幻灯片；文字和图形不可单独编辑。')


def _pdf_to_word(source, target, options, cancel, progress, notes):
    with _open_pdf(source) as document:
        if options.pdf_word_mode == 'image':
            from docx import Document
            from docx.enum.section import WD_SECTION
            from docx.shared import Inches, Pt
            word = Document()
            for index, page in enumerate(document):
                check_cancel(cancel)
                section = word.sections[0] if index == 0 else word.add_section(WD_SECTION.NEW_PAGE)
                section.page_width = Inches(page.rect.width / 72)
                section.page_height = Inches(page.rect.height / 72)
                section.top_margin = section.bottom_margin = Inches(0.08)
                section.left_margin = section.right_margin = Inches(0.08)
                section.header_distance = section.footer_distance = 0
                paragraph = word.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                available_w = page.rect.width / 72 - 0.16
                available_h = page.rect.height / 72 - 0.35
                if min(available_w, available_h) <= 0:
                    raise ConversionError('PDF 页面尺寸过小，无法放入 Word。')
                scale = min(available_w / page.rect.width, available_h / page.rect.height)
                pixmap = page.get_pixmap(dpi=options.dpi, alpha=False)
                paragraph.add_run().add_picture(io.BytesIO(pixmap.tobytes('png')),
                                                width=Inches(page.rect.width * scale),
                                                height=Inches(page.rect.height * scale))
                progress((index + 1) / len(document) * 0.95, f'转换 PDF 页 {index + 1}/{len(document)}')
            word.save(target)
            notes.append('整页图片模式：保留页面外观，文字不可编辑，未进行 OCR。')
            return
        scanned = [i + 1 for i, page in enumerate(document) if not page.get_text().strip()]
        if len(scanned) == len(document):
            raise ConversionError('此 PDF 没有可提取的文字。请改用「整页图片」模式；可编辑扫描件需要 OCR。')
        if scanned:
            notes.append(f'有 {len(scanned)} 页未检测到文字；这些页面不会获得可编辑文字。')
    from pdf2docx import Converter
    progress(0.1, '恢复 Word 文字和布局；取消将在当前文件结束后生效')
    check_cancel(cancel)
    converter = Converter(str(source))
    try:
        converter.convert(str(target), start=0, end=None, multi_processing=False)
    finally:
        converter.close()
    notes.append('可编辑模式尽力恢复文字和表格；复杂布局、公式、字体可能需要调整。')


def _office_pdf(source, target, options, cancel, progress, notes):
    is_word = source.suffix.lower()[1:] in WORD_INPUTS
    errors = []
    check_cancel(cancel)
    if options.office_backend in {'auto', 'office'} and sys.platform == 'win32':
        app = document = None
        initialized = False
        owns_application = False
        original_settings = {}
        try:
            import pythoncom
            import win32com.client
            process_name = 'WINWORD.EXE' if is_word else 'POWERPNT.EXE'
            snapshot = subprocess.run(['tasklist', '/FI', f'IMAGENAME eq {process_name}', '/FO', 'CSV', '/NH'],
                                      capture_output=True, timeout=15, creationflags=CREATE_NO_WINDOW)
            if snapshot.returncode:
                raise ConversionError('无法检查 Office 进程，请使用 LibreOffice 引擎。')
            existing = set()
            for row in csv.reader(snapshot.stdout.decode(errors='replace').splitlines()):
                if len(row) > 1 and row[1].isdigit():
                    existing.add(int(row[1]))
            pythoncom.CoInitialize()
            initialized = True
            progress(0.15, '使用本机 Office 导出 PDF；取消将在当前文件结束后生效')
            app = win32com.client.DispatchEx('Word.Application' if is_word else 'PowerPoint.Application')
            # Word supports an independent DispatchEx application. PowerPoint
            # can reuse its singleton, so never quit it if it was already running.
            owns_application = is_word or not existing
            for setting in ('AutomationSecurity', 'DisplayAlerts'):
                value = getattr(app, setting)
                original_settings[setting] = value() if callable(value) else value
            app.AutomationSecurity = 3  # Disable macros in documents opened by automation.
            # PowerPoint can reuse an existing application. Opening a private copy
            # prevents closing an original presentation the user already has open.
            private_source = target.parent / 'office-input' / source.name
            private_source.parent.mkdir()
            shutil.copy2(source, private_source)
            if is_word:
                if owns_application:
                    app.Visible = False
                app.DisplayAlerts = 0
                document = app.Documents.Open(str(private_source), ConfirmConversions=False, ReadOnly=True,
                                              AddToRecentFiles=False, PasswordDocument='__NO_PASSWORD__')
                document.ExportAsFixedFormat(str(target), 17)
            else:
                app.DisplayAlerts = 1
                document = app.Presentations.Open(str(private_source), ReadOnly=True, Untitled=False, WithWindow=False)
                document.SaveAs(str(target), 32)
            if not target.exists() or not target.stat().st_size:
                raise ConversionError('Office 没有生成有效的 PDF。')
            notes.append('使用本机 Microsoft Office 导出 PDF。')
            return
        except Exception as exc:
            errors.append(f'Microsoft Office：{exc}')
        finally:
            if document is not None:
                try:
                    document.Close(False) if is_word else document.Close()
                except Exception:
                    pass
            if app is not None:
                for setting, value in original_settings.items():
                    try:
                        setattr(app, setting, value)
                    except Exception:
                        pass
                try:
                    if owns_application:
                        app.Quit()
                except Exception:
                    pass
            if initialized:
                pythoncom.CoUninitialize()
    check_cancel(cancel)
    target.unlink(missing_ok=True)
    executable = libreoffice_path() if options.office_backend in {'auto', 'libreoffice'} else None
    if executable:
        progress(0.2, '使用 LibreOffice 导出 PDF')
        with tempfile.TemporaryDirectory(prefix='office-profile-') as profile:
            filter_name = 'writer_pdf_Export' if is_word else 'impress_pdf_Export'
            command = [executable, f'-env:UserInstallation={Path(profile).as_uri()}', '--headless',
                       '--convert-to', f'pdf:{filter_name}', '--outdir', str(target.parent), str(source)]
            process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                       creationflags=CREATE_NO_WINDOW)
            started = time.monotonic()
            while process.poll() is None:
                if cancel.is_set() or time.monotonic() - started > options.timeout:
                    process.terminate()
                    try:
                        process.wait(timeout=5)
                    except subprocess.TimeoutExpired:
                        process.kill()
                        process.wait()
                    if cancel.is_set():
                        raise Cancelled('已取消')
                    raise ConversionError('LibreOffice 转换超时。')
                time.sleep(0.1)
            output = process.communicate()[0].decode('utf-8', errors='replace')
            produced = target.parent / (source.stem + '.pdf')
            if process.returncode == 0 and produced.is_file() and produced.stat().st_size:
                if produced != target:
                    produced.replace(target)
                notes.append('使用 LibreOffice 导出 PDF。')
                return
            errors.append(f'LibreOffice：{output[-1200:]}')
    detail = '\n'.join(errors)
    raise ConversionError('Office 导出失败。需要已激活的 Word/PowerPoint，或安装 LibreOffice。'
                          '\n请检查文件是否损坏、加密或正被其他程序锁定。' + ('\n' + detail if detail else ''))


def _duration(executable: str, source: Path) -> float:
    try:
        completed = subprocess.run([executable, '-hide_banner', '-i', str(source)],
                                   stdin=subprocess.DEVNULL, capture_output=True, timeout=20,
                                   creationflags=CREATE_NO_WINDOW)
        match = re.search(r'Duration: (\d+):(\d+):(\d+\.\d+)', completed.stderr.decode(errors='replace'))
        if match:
            return int(match[1]) * 3600 + int(match[2]) * 60 + float(match[3])
    except (OSError, subprocess.TimeoutExpired):
        pass
    return 0


def _video(source, target, options, cancel, progress, notes):
    from media import video_filters, media_notes
    filters = video_filters(options)
    notes.extend(media_notes(options))
    executable = ffmpeg_path()
    duration = _duration(executable, source)
    extension = target.suffix[1:]
    crf = {'small': 28, 'balanced': 23, 'high': 18}[options.video_quality]
    command = [executable, '-y', '-hide_banner', '-nostdin', '-loglevel', 'warning']
    if source.suffix.lower() == '.gif':
        command += ['-ignore_loop', '1']
    command += ['-i', str(source)]
    if extension == 'gif':
        prefix = ','.join(filters) + ',' if filters else ''
        command += ['-filter_complex', prefix + 'fps=15,scale=min(960\\,iw):-1:flags=lanczos,split[a][b];'
                    '[a]palettegen[p];[b][p]paletteuse', '-an', '-loop', '0']
        notes.append('GIF 输出为 15 帧/秒，最大宽度 960 像素，无音频。')
    else:
        filters.append('pad=ceil(iw/2)*2:ceil(ih/2)*2')
        command += ['-map', '0:v:0', '-map', '0:a:0?', '-vf', ','.join(filters), '-pix_fmt', 'yuv420p']
        if source.suffix.lower() == '.gif':
            # GIF delays use centiseconds. A 100-fps timeline retains the last
            # frame's duration, which default FFmpeg frame-rate inference loses.
            command += ['-r', '100', '-fps_mode', 'cfr']
        if extension == 'webm':
            command += ['-c:v', 'libvpx-vp9', '-crf', str(crf + 8), '-b:v', '0', '-c:a', 'libopus', '-b:a', '128k']
        elif extension == 'avi':
            command += ['-c:v', 'mpeg4', '-q:v', '3', '-c:a', 'libmp3lame', '-b:a', '192k']
        elif extension == 'wmv':
            command += ['-c:v', 'wmv2', '-b:v', '4M', '-c:a', 'wmav2', '-b:a', '128k']
        elif extension == 'mpg':
            command += ['-c:v', 'mpeg2video', '-q:v', '3', '-r', '25', '-c:a', 'mp2', '-b:a', '192k', '-ar', '48000']
        elif extension == 'flv':
            command += ['-c:v', 'libx264', '-crf', str(crf), '-preset', 'medium', '-c:a', 'aac', '-ar', '44100']
        else:
            command += ['-c:v', 'libx264', '-crf', str(crf), '-preset', 'medium', '-c:a', 'aac', '-b:a', '192k']
            if extension in {'mp4', 'mov', 'm4v'}:
                command += ['-movflags', '+faststart']
        notes.append('视频重新编码，保留第一条视频和第一条音轨；字幕及其他音轨不会写入。')
    command += ['-progress', 'pipe:1', '-nostats', str(target)]
    lines = queue.Queue()
    with tempfile.TemporaryFile() as error_log:
        process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=error_log,
                                   stdin=subprocess.DEVNULL, creationflags=CREATE_NO_WINDOW)
        def collect():
            for line in iter(process.stdout.readline, b''):
                lines.put(line.decode('utf-8', errors='replace').strip())
            process.stdout.close()
        reader = threading.Thread(target=collect, daemon=True)
        reader.start()
        started = time.monotonic()
        try:
            while process.poll() is None:
                check_cancel(cancel)
                if time.monotonic() - started > options.timeout:
                    raise ConversionError(f'视频转换超过 {options.timeout} 秒，请提高超时设置。')
                try:
                    line = lines.get(timeout=0.1)
                except queue.Empty:
                    continue
                if line.startswith('out_time_us='):
                    try:
                        seconds = int(line.partition('=')[2]) / 1_000_000
                        percent = min(0.98, seconds / duration) if duration else 0
                        progress(percent, f'已转换 {seconds:.1f} 秒')
                    except ValueError:
                        pass
            reader.join(timeout=2)
            check_cancel(cancel)
            if process.returncode:
                error_log.seek(0, 2)
                size = error_log.tell()
                error_log.seek(max(0, size - 5000))
                detail = error_log.read().decode('utf-8', errors='replace')
                raise ConversionError('FFmpeg 转换失败：\n' + detail)
        finally:
            if process.poll() is None:
                process.terminate()
                try:
                    process.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait()
            reader.join(timeout=2)


def convert(source: Path | str, directory: Path | str, extension: str, category: str,
            options: Options | None = None, cancel: threading.Event | None = None,
            progress: Callable[[float, str], None] | None = None) -> tuple[Path, list[str]]:
    source = Path(source).resolve()
    directory = Path(directory).resolve()
    options = options or Options()
    cancel = cancel or threading.Event()
    progress = progress or (lambda fraction, message: None)
    extension = extension.lower().lstrip('.')
    if not source.is_file():
        raise ConversionError('源文件不存在：' + str(source))
    if extension not in targets_for(source, category):
        raise ConversionError(f'不支持此转换：{source.suffix} → {extension}（{category}）')
    if not 1 <= options.quality <= 100 or not 72 <= options.dpi <= 600:
        raise ConversionError('图片质量应为 1–100，PDF 分辨率应为 72–600 DPI。')
    if options.pdf_word_mode not in {'editable', 'image'} or options.video_quality not in {'small', 'balanced', 'high'}:
        raise ConversionError('无效的转换设置。')
    if options.office_backend not in {'auto', 'office', 'libreoffice'} or options.timeout < 1:
        raise ConversionError('无效的引擎或超时设置。')
    from media import validate_media
    validate_media(options)
    directory.mkdir(parents=True, exist_ok=True)
    check_cancel(cancel)
    notes = []
    with tempfile.TemporaryDirectory(prefix='.convert-', dir=directory) as temporary:
        target = Path(temporary) / f'result.{extension}'
        ext = source.suffix.lower()[1:]
        if category == CATEGORIES[2]:
            _image(source, target, options, cancel, progress, notes)
        elif ext == 'pdf' and extension == 'docx':
            _pdf_to_word(source, target, options, cancel, progress, notes)
        elif ext == 'pdf' and extension == 'pptx':
            _pdf_to_ppt(source, target, options, cancel, progress, notes)
        elif extension == 'pdf' and ext in WORD_INPUTS | PPT_INPUTS:
            _office_pdf(source, target, options, cancel, progress, notes)
        else:
            _video(source, target, options, cancel, progress, notes)
        check_cancel(cancel)
        if not target.is_file() or target.stat().st_size == 0:
            raise ConversionError('转换引擎未生成有效文件。')
        progress(0.99, '正在保存结果')
        destination = publish(target, directory, source.stem, extension)
    progress(1.0, '转换完成')
    return destination, notes
