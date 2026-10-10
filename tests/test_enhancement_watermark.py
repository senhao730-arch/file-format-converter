import io
import json
import subprocess
import threading
import time
import zipfile
from pathlib import Path

import numpy as np
import pikepdf
import pymupdf
import pytest
from PIL import Image, ImageDraw

from converter import Options, ConversionError, Cancelled, convert, ffmpeg_path
from media import adjust_image
import watermark as wm


@pytest.mark.parametrize('mode', ['contain','crop','stretch'])
def test_image_resolution_and_alpha(mode):
    image=Image.new('RGBA',(120,60),(255,30,20,150))
    output=adjust_image(image,Options(width=80,height=80,resize_mode=mode,sharpen=1,denoise=1))
    assert output.size==(80,80)
    assert output.getpixel((40,40))[3]==150
    if mode=='contain':
        assert output.getpixel((0,0))[3]==0
    assert adjust_image(image,Options(height=30)).size==(60,30)
    with pytest.raises(ConversionError):
        adjust_image(image,Options(width=-1))


def image_fixture(path):
    image=Image.new('RGBA',(300,200),(235,240,245,190))
    ImageDraw.Draw(image).text((205,175),'WATERMARK',fill=(40,40,40,190))
    image.save(path)
    return image


def test_image_detect_repair_crop_no_overwrite_and_fingerprint(tmp_path):
    source=tmp_path/'picture.png'
    original=image_fixture(source)
    digest=wm.fingerprint(source)
    analysis=wm.detect(source)
    assert analysis.candidates
    region=wm.add_region(analysis,0,(.65,.83,.35,.17))
    result,_=wm.remove(analysis,[region.id],tmp_path/'out')
    repaired=Image.open(result)
    assert repaired.size==original.size
    assert np.array_equal(np.array(repaired)[:150],np.array(original)[:150])
    assert np.array_equal(np.array(repaired)[:,:,3],np.array(original)[:,:,3])
    patch=np.array(repaired)[170:190,210:280,:3]
    before=np.array(original)[170:190,210:280,:3]
    assert np.mean(abs(patch.astype(float)-235))<np.mean(abs(before.astype(float)-235))
    again,_=wm.remove(analysis,[region.id],tmp_path/'out')
    assert again!=result and result.is_file()
    crop,_=wm.remove(analysis,[region.id],tmp_path/'out','crop')
    assert Image.open(crop).height<original.height
    assert wm.fingerprint(source)==digest
    source.write_bytes(b'changed')
    with pytest.raises(ConversionError,match='变化'):
        wm.remove(analysis,[region.id],tmp_path/'out')


def test_animated_repair_and_cancel(tmp_path):
    source=tmp_path/'animation.gif'
    frames=[Image.new('RGB',(80,60),color) for color in ('red','blue')]
    frames[0].save(source,save_all=True,append_images=frames[1:],duration=[100,200],loop=2)
    analysis=wm.detect(source)
    region=wm.add_region(analysis,0,(.1,.1,.2,.2))
    result,_=wm.remove(analysis,[region.id],tmp_path/'out')
    with Image.open(result) as image:
        assert image.n_frames==2 and image.info['loop']==2
        image.seek(1)
        assert image.info['duration']==200
    cancelled=threading.Event()
    cancelled.set()
    with pytest.raises(Cancelled):
        wm.remove(analysis,[region.id],tmp_path/'out',cancel=cancelled)


def pdf_fixture(path,rotation=0):
    pdf=pymupdf.open()
    for index in range(2):
        page=pdf.new_page(width=300,height=200)
        page.insert_text((30,50),f'BODY PAGE {index+1}')
        page.set_rotation(rotation)
    pdf.save(path)
    pdf.close()
    with pikepdf.Pdf.open(path,allow_overwriting_input=True) as document:
        form=document.make_stream(b'0.9 0.1 0.1 rg 0 0 50 15 re f')
        form['/Type']=pikepdf.Name('/XObject')
        form['/Subtype']=pikepdf.Name('/Form')
        form['/BBox']=pikepdf.Array([0,0,50,15])
        form['/Resources']=pikepdf.Dictionary()
        for page in document.pages:
            resources=pikepdf.Dictionary(page.Resources)
            resources['/XObject']=pikepdf.Dictionary(WMlogo=form)
            page.Resources=resources
            original=page.Contents
            streams=list(original) if isinstance(original,pikepdf.Array) else [original]
            page.Contents=pikepdf.Array(streams+[document.make_stream(b'q 1 0 0 1 220 20 cm /WMlogo Do Q')])
        document.save(path)


def test_pdf_object_removal_is_page_scoped_and_preserves_text(tmp_path):
    source=tmp_path/'native.pdf'
    pdf_fixture(source)
    digest=wm.fingerprint(source)
    analysis=wm.detect(source)
    candidates=[c for c in analysis.candidates if c.kind=='pdf_object']
    assert len(candidates)==2 and all(c.box for c in candidates)
    first=candidates[0]
    # Form is at PDF y=20 (bottom); displayed y is 165 on the unrotated page.
    assert first.box[0]==pytest.approx(220/300)
    assert first.box[1]==pytest.approx(165/200)
    result,_=wm.remove(analysis,[first.id],tmp_path/'out')
    with pymupdf.open(source) as before,pymupdf.open(result) as after:
        assert 'BODY PAGE 1' in after[0].get_text()
        assert 'BODY PAGE 2' in after[1].get_text()
        assert before[1].get_pixmap().samples==after[1].get_pixmap().samples
        assert before[0].get_pixmap().samples!=after[0].get_pixmap().samples
    assert wm.fingerprint(source)==digest


def test_pdf_text_candidate_and_stamp_annotation(tmp_path):
    source=tmp_path/'text_stamp.pdf'
    with pymupdf.open() as pdf:
        page=pdf.new_page(width=400,height=300)
        page.insert_text((40,40),'BODY TEXT')
        point=pymupdf.Point(50,230)
        page.insert_text(point,'CONFIDENTIAL',fontsize=28,color=(.8,.8,.8),
                         morph=(point,pymupdf.Matrix(1,1).prerotate(30)))
        page.add_stamp_annot(pymupdf.Rect(250,240,370,275))
        pdf.save(source)
    analysis=wm.detect(source)
    assert any(c.kind=='region' for c in analysis.candidates)
    annotation=next(c for c in analysis.candidates if c.kind=='pdf_annotation')
    result,_=wm.remove(analysis,[annotation.id],tmp_path/'out')
    with pymupdf.open(result) as pdf:
        assert not list(pdf[0].annots() or [])
        assert 'BODY TEXT' in pdf[0].get_text()
        assert 'CONFIDENTIAL' in pdf[0].get_text()


@pytest.mark.parametrize('rotation',[0,90,180,270])
def test_pdf_region_repair_rotations_only_selected_page_loses_text(tmp_path,rotation):
    source=tmp_path/'rotated.pdf'
    pdf_fixture(source,rotation)
    analysis=wm.detect(source)
    object_candidate=next(c for c in analysis.candidates if c.kind=='pdf_object' and c.page==0)
    region=wm.add_region(analysis,0,object_candidate.box)
    result,notes=wm.remove(analysis,[region.id],tmp_path/'out')
    assert any('文字' in note for note in notes)
    with pymupdf.open(source) as before,pymupdf.open(result) as after:
        assert not after[0].get_text().strip()
        assert 'BODY PAGE 2' in after[1].get_text()
        assert after[0].rotation==rotation
        assert after[0].rect==before[0].rect
        old=np.asarray(wm.preview(analysis,0)).astype(float)
        new=np.asarray(wm.preview(analysis,0,result)).astype(float)
        box=wm.pixel_boxes([region],old.shape[1],old.shape[0])[0]
        x,y,w,h=box
        # The red watermark must be repaired in the selected display area.
        red_before=((old[y:y+h,x:x+w,0]>150)&(old[y:y+h,x:x+w,1]<100)).sum()
        red_after=((new[y:y+h,x:x+w,0]>150)&(new[y:y+h,x:x+w,1]<100)).sum()
        assert red_before>0 and red_after<red_before*.5
        # Outside the repair region the full rendered page is still aligned.
        mask=np.ones(old.shape[:2],bool)
        mask[max(0,y-3):y+h+3,max(0,x-3):x+w+3]=False
        assert np.mean(abs(old[mask]-new[mask]))<4


def word_fixture(path):
    from docx import Document
    from docx.oxml import parse_xml
    document=Document()
    document.add_paragraph('正文必须保留 BODY')
    header=document.sections[0].header.paragraphs[0]
    header.text='页眉必须保留 HEADER'
    for index in range(2):
        header._p.append(parse_xml(f'<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:shape id="PowerPlusWaterMarkObject{index}" style="position:absolute"><v:textpath string="TEST WATERMARK {index}"/></v:shape></w:pict>'))
    document.save(path)


def test_word_selected_shapes_preserves_header_body_and_zip_parts(tmp_path):
    from docx import Document
    source=tmp_path/'watermark.docx'
    word_fixture(source)
    analysis=wm.detect(source)
    assert len(analysis.candidates)==2
    result,_=wm.remove(analysis,[analysis.candidates[0].id],tmp_path/'out')
    document=Document(result)
    assert document.paragraphs[0].text=='正文必须保留 BODY'
    assert document.sections[0].header.paragraphs[0].text=='页眉必须保留 HEADER'
    remaining=wm.detect(result)
    assert len(remaining.candidates)==1
    assert remaining.candidates[0].label=='TEST WATERMARK 1'
    with zipfile.ZipFile(source) as before,zipfile.ZipFile(result) as after:
        assert before.namelist()==after.namelist()
        assert before.read('word/document.xml')==after.read('word/document.xml')
    both,_=wm.remove(analysis,[c.id for c in analysis.candidates],tmp_path/'out')
    assert not wm.detect(both).candidates


def test_word_child_namespace_locators(tmp_path):
    source=tmp_path/'namespaces.docx'
    xml=b'<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml" id="PowerPlusWaterMarkObject1"><v:textpath string="WATERMARK"/></v:shape></w:pict></w:p></w:hdr>'
    with zipfile.ZipFile(source,'w') as archive:
        archive.writestr('word/header1.xml',xml)
    analysis=wm.detect(source)
    assert len(analysis.candidates)==1
    result,_=wm.remove(analysis,[analysis.candidates[0].id],tmp_path/'out')
    assert not wm.detect(result).candidates


def test_video_resolution_audio_and_watermark_timing(tmp_path):
    source=tmp_path/'video.mp4'
    command=[ffmpeg_path(),'-hide_banner','-loglevel','error','-f','lavfi','-i','color=c=white:s=160x100:r=10:d=1.5',
             '-f','lavfi','-i','sine=frequency=440:duration=1.5','-vf','drawbox=x=115:y=75:w=30:h=12:color=red:t=fill',
             '-c:v','libx264','-pix_fmt','yuv420p','-c:a','aac','-shortest',str(source)]
    subprocess.run(command,check=True,capture_output=True)
    resized,_=convert(source,tmp_path/'out','mp4','视频',Options(width=81,height=61,sharpen=1,denoise=1))
    assert wm._frame(resized).size==(82,62)
    analysis=wm.detect(source)
    assert len(analysis.frames)==3 and analysis.width==160
    region=wm.add_region(analysis,0,(110/160,70/100,40/160,22/100))
    result,_=wm.remove(analysis,[region.id],tmp_path/'out',start=.5,end=1.2)
    before=np.array(wm._frame(result,.1))
    during=np.array(wm._frame(result,.8))
    assert before[80,125,0]>180 and before[80,125,1]<80
    assert during[80,125,1]>180
    audio=subprocess.run([ffmpeg_path(),'-hide_banner','-i',str(result)],capture_output=True).stderr.decode('utf-8','replace')
    assert 'Audio:' in audio


def exercise_watermark_gui(tmp_path):
    import tkinter as tk
    from types import SimpleNamespace
    from watermark_ui import WatermarkWindow
    source=tmp_path/'image.png'
    image_fixture(source)
    root=tk.Tk()
    root.geometry('1x1+0+0')
    window=WatermarkWindow(root,str(tmp_path/'out'))
    try:
        window.source.set(str(source))
        window.detect()
        deadline=time.time()+20
        while window.busy and time.time()<deadline:
            root.update()
            time.sleep(.02)
        assert window.analysis and not window.selected
        root.update()
        x,y,w,h=window.display_rect
        window.drag_begin(SimpleNamespace(x=x+w*.65,y=y+h*.83))
        window.drag_end(SimpleNamespace(x=x+w,y=y+h))
        candidate=window.analysis.candidates[-1]
        assert candidate.box==pytest.approx((.65,.83,.35,.17))
        assert candidate.id in window.selected
        window.process(True)
        deadline=time.time()+20
        while window.busy and time.time()<deadline:
            root.update()
            time.sleep(.02)
        assert window.result.is_file() and window.view.get()=='处理效果'
        window.clear_selection()
        assert window.result is None
    finally:
        window.close()
        root.destroy()


def test_watermark_gui_manual_selection_and_preview(tmp_path):
    # Isolate Tcl interpreters and pending GUI callbacks from other GUI tests.
    import sys
    command=[sys.executable,'-c',
             'import runpy,sys; from pathlib import Path; runpy.run_path(sys.argv[1])["exercise_watermark_gui"](Path(sys.argv[2]))',
             str(Path(__file__).resolve()),str(tmp_path)]
    completed=subprocess.run(command,capture_output=True,timeout=60)
    assert completed.returncode==0,completed.stderr.decode('utf-8','replace')
