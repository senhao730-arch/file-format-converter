"""Integration tests use actual generated files rather than mocked conversions."""
import os
import subprocess
import threading
from pathlib import Path

import pytest
from PIL import Image
import pymupdf
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from converter import CATEGORIES, IMAGE_OUTPUTS, VIDEO_OUTPUTS, Options, Cancelled, ConversionError, convert, ffmpeg_path


@pytest.fixture
def sample_pdf(tmp_path):
    path = tmp_path / '中文 文件.pdf'
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), 'Editable text for conversion.', fontsize=18)
    page.draw_rect((72, 150, 300, 250), fill=(0.2, 0.5, 0.9))
    page = doc.new_page(width=842, height=595)
    page.insert_text((72, 100), 'Second page, landscape.', fontsize=16)
    doc.save(path)
    doc.close()
    return path


@pytest.mark.parametrize('extension', IMAGE_OUTPUTS)
def test_images(tmp_path, extension):
    source = tmp_path / '透明 图片.png'
    Image.new('RGBA', (63, 47), (255, 0, 0, 0)).save(source)
    result, notes = convert(source, tmp_path / 'out', extension, CATEGORIES[2])
    with Image.open(result) as im:
        im.load()
        assert im.width > 0 and im.height > 0
        if extension == 'jpg':
            assert min(im.convert('RGB').getpixel((0, 0))) >= 250
    assert source.exists()


@pytest.mark.parametrize('extension', ['gif', 'webp', 'png', 'tiff', 'avif'])
def test_animation_retains_frames(tmp_path, extension):
    source = tmp_path / '动画.gif'
    frames = [Image.new('RGB', (65, 49), color) for color in ('red', 'green', 'blue')]
    frames[0].save(source, save_all=True, append_images=frames[1:], duration=[100, 200, 300], loop=0)
    result, notes = convert(source, tmp_path / 'out', extension, CATEGORIES[2])
    with Image.open(result) as im:
        assert im.n_frames == 3
        if extension in {'gif', 'webp', 'png'}:
            durations = []
            for index in range(im.n_frames):
                im.seek(index)
                im.load()
                durations.append(im.info.get('duration', 0))
            assert abs(sum(durations) - 600) < 50


def test_collision_does_not_overwrite_source_or_existing(tmp_path):
    source = tmp_path / 'a.png'
    Image.new('RGB', (10, 10), 'red').save(source)
    before = source.read_bytes()
    first, _ = convert(source, tmp_path, 'png', CATEGORIES[2])
    second, _ = convert(source, tmp_path, 'png', CATEGORIES[2])
    assert source.read_bytes() == before
    assert first.name == 'a (1).png'
    assert second.name == 'a (2).png'


def test_pdf_to_editable_word(tmp_path, sample_pdf):
    result, _ = convert(sample_pdf, tmp_path / 'out', 'docx', CATEGORIES[0])
    document = Document(result)
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'Editable text' in text
    assert 'Second page' in text


def test_pdf_to_image_word(tmp_path, sample_pdf):
    result, _ = convert(sample_pdf, tmp_path / 'out', 'docx', CATEGORIES[0], Options(pdf_word_mode='image'))
    document = Document(result)
    assert len(document.inline_shapes) == 2
    assert len(document.sections) == 2


def test_pdf_to_ppt(tmp_path, sample_pdf):
    result, _ = convert(sample_pdf, tmp_path / 'out', 'pptx', CATEGORIES[1])
    presentation = Presentation(result)
    assert len(presentation.slides) == 2
    assert all(len(slide.shapes) == 1 for slide in presentation.slides)


def test_scanned_pdf_modes(tmp_path):
    source = tmp_path / '扫描.pdf'
    doc = pymupdf.open()
    page = doc.new_page()
    import io
    image = io.BytesIO()
    Image.new('RGB', (100, 100), 'blue').save(image, format='PNG')
    page.insert_image(page.rect, stream=image.getvalue())
    doc.save(source)
    doc.close()
    with pytest.raises(ConversionError, match='OCR'):
        convert(source, tmp_path / 'out', 'docx', CATEGORIES[0])
    result, _ = convert(source, tmp_path / 'out', 'docx', CATEGORIES[0], Options(pdf_word_mode='image'))
    assert len(Document(result).inline_shapes) == 1


def test_password_pdf(tmp_path):
    source = tmp_path / 'locked.pdf'
    doc = pymupdf.open()
    doc.new_page()
    doc.save(source, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='secret')
    doc.close()
    with pytest.raises(ConversionError, match='密码'):
        convert(source, tmp_path / 'out', 'pptx', CATEGORIES[1])


def test_cancellation_and_corrupt_cleanup(tmp_path):
    source = tmp_path / 'bad.png'
    source.write_bytes(b'not an image')
    output = tmp_path / 'out'
    event = threading.Event()
    event.set()
    with pytest.raises(Cancelled):
        convert(source, output, 'jpg', CATEGORIES[2], cancel=event)
    with pytest.raises(Exception):
        convert(source, output, 'jpg', CATEGORIES[2])
    assert list(output.iterdir()) == []


@pytest.fixture
def sample_video(tmp_path):
    source = tmp_path / '有声音 视频.mp4'
    subprocess.run([ffmpeg_path(), '-hide_banner', '-loglevel', 'error', '-f', 'lavfi', '-i',
                    'testsrc2=size=64x48:rate=10:duration=0.8', '-f', 'lavfi', '-i',
                    'sine=frequency=440:sample_rate=48000:duration=0.8', '-c:v', 'libx264',
                    '-pix_fmt', 'yuv420p', '-c:a', 'aac', '-shortest', str(source)], check=True, capture_output=True)
    return source


@pytest.mark.parametrize('extension', VIDEO_OUTPUTS)
def test_video_formats(tmp_path, sample_video, extension):
    result, _ = convert(sample_video, tmp_path / 'out', extension, CATEGORIES[4])
    probe = subprocess.run([ffmpeg_path(), '-hide_banner', '-i', str(result), '-f', 'null', '-'], capture_output=True)
    assert probe.returncode == 0, probe.stderr.decode(errors='replace')
    if extension != 'gif':
        assert b'Audio:' in probe.stderr


def test_gif_to_mp4_odd_dimensions(tmp_path):
    source = tmp_path / '奇数尺寸.gif'
    frames = [Image.new('RGB', (65, 49), color) for color in ('red', 'green', 'blue')]
    frames[0].save(source, save_all=True, append_images=frames[1:], duration=[100, 200, 300], loop=0)
    result, _ = convert(source, tmp_path / 'out', 'mp4', CATEGORIES[3])
    decoded = subprocess.run([ffmpeg_path(), '-hide_banner', '-i', str(result), '-f', 'null', '-'], capture_output=True)
    assert decoded.returncode == 0
    assert b'66x50' in decoded.stderr
    assert b'00:00:00.60' in decoded.stderr


def test_video_cancellation(tmp_path):
    source = tmp_path / 'cancel.gif'
    frames = [Image.new('RGB', (64, 48), color) for color in ('red', 'blue')]
    frames[0].save(source, save_all=True, append_images=frames[1:], duration=100)
    event = threading.Event()
    def stop(fraction, message):
        event.set()
    with pytest.raises(Cancelled):
        convert(source, tmp_path / 'out', 'mp4', CATEGORIES[3], cancel=event, progress=stop)
    assert list((tmp_path / 'out').iterdir()) == []


@pytest.mark.skipif(os.environ.get('RUN_OFFICE_TESTS') != '1', reason='Requires activated Microsoft Office')
def test_office_roundtrip(tmp_path, sample_pdf):
    word = tmp_path / '示例.docx'
    document = Document()
    document.add_heading('Office conversion test', 0)
    document.add_paragraph('Word to PDF works.')
    document.save(word)
    pdf, _ = convert(word, tmp_path / 'out', 'pdf', CATEGORIES[0], Options(office_backend='office'))
    with pymupdf.open(pdf) as doc:
        assert 'Word to PDF works.' in doc[0].get_text()
    ppt = tmp_path / '示例.pptx'
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = 'PowerPoint to PDF works.'
    presentation.save(ppt)
    pdf, _ = convert(ppt, tmp_path / 'out', 'pdf', CATEGORIES[1], Options(office_backend='office'))
    with pymupdf.open(pdf) as doc:
        assert 'PowerPoint to PDF works.' in doc[0].get_text()
    visual_word, _ = convert(sample_pdf, tmp_path / 'out', 'docx', CATEGORIES[0], Options(pdf_word_mode='image'))
    rendered, _ = convert(visual_word, tmp_path / 'roundtrip', 'pdf', CATEGORIES[0], Options(office_backend='office'))
    with pymupdf.open(rendered) as doc:
        assert len(doc) == 2, 'Image Word unexpectedly produced extra pages'
